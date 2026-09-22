"""Turn a spreadsheet into a table.

Reads CSV (and tab-separated files) and Excel `.xlsx` workbooks. Cells are
only ever read as data. In a CSV a formula is just the text it is. In an
Excel file the value Excel last worked out is used; the formula itself is
never run by Omarchy-DB.
"""

from __future__ import annotations

import csv
import datetime as _dt
from collections.abc import Sequence
from pathlib import Path
from typing import Any

from .errors import ImportProblem
from .fields import Field, coerce, slugify_name
from .infer import infer_fields
from .paths import resolve_under_roots
from .storage.base import Storage

SAMPLE_ROWS = 500
CSV_SUFFIXES = {".csv", ".tsv", ".txt"}
EXCEL_SUFFIXES = {".xlsx", ".xlsm"}
OLD_EXCEL_SUFFIXES = {".xls"}


def read_csv(path: Path, *, max_rows: int | None = None) -> tuple[list[str], list[list[str]]]:
    """Read a CSV file into headings plus rows of plain text."""
    with path.open("r", newline="", encoding="utf-8-sig", errors="replace") as handle:
        sample = handle.read(64 * 1024)
        handle.seek(0)
        try:
            dialect: Any = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(handle, dialect)
        try:
            headers = next(reader)
        except StopIteration:
            raise ImportProblem("That file is empty.") from None
        rows: list[list[str]] = []
        for row in reader:
            if not any(str(cell).strip() for cell in row):
                continue
            rows.append([str(cell) for cell in row])
            if max_rows is not None and len(rows) >= max_rows:
                break
    headers = [str(header) for header in headers]
    if not any(header.strip() for header in headers):
        raise ImportProblem("The first line of that file has no column names.")
    return headers, rows


def _openpyxl():
    try:
        import openpyxl  # noqa: PLC0415
    except ImportError as exc:  # pragma: no cover - depends on the machine
        raise ImportProblem(
            "Reading Excel files needs the 'openpyxl' package. "
            "Install it with: pip install 'omarchy-db[xlsx]'"
        ) from exc
    return openpyxl


def _excel_text(value: Any) -> str:
    """Excel cells come typed. Turn each into the text the type guesser reads."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "yes" if value else "no"
    if isinstance(value, _dt.datetime):
        if value.time() == _dt.time(0, 0):
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, _dt.date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _load_book(path: Path):
    openpyxl = _openpyxl()
    try:
        return openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises many kinds
        raise ImportProblem(f"That does not look like an Excel file: {exc}") from exc


def list_sheets(path: Path) -> list[str]:
    """The sheet names in a workbook; a CSV has one unnamed sheet."""
    if path.suffix.lower() not in EXCEL_SUFFIXES:
        return []
    book = _load_book(path)
    try:
        return list(book.sheetnames)
    finally:
        book.close()


def read_xlsx(
    path: Path, *, sheet: str | None = None, max_rows: int | None = None
) -> tuple[list[str], list[list[str]]]:
    """Read one sheet of an Excel workbook into headings plus rows of plain text.

    `data_only=True` means a formula cell gives the value Excel last saved for
    it. Nothing is calculated here, and no macro is ever touched.
    """
    book = _load_book(path)
    try:
        if sheet:
            if sheet not in book.sheetnames:
                names = ", ".join(book.sheetnames)
                raise ImportProblem(f"There is no sheet called {sheet!r}. This file has: {names}.")
            worksheet = book[sheet]
        else:
            worksheet = book.worksheets[0]
        rows_iter = worksheet.iter_rows(values_only=True)
        headers: list[str] = []
        for raw in rows_iter:
            headers = [_excel_text(cell).strip() for cell in raw]
            if any(headers):
                break
        else:
            raise ImportProblem("That sheet is empty.")
        # Drop trailing blank heading columns Excel likes to leave behind.
        while headers and not headers[-1]:
            headers.pop()
        if not headers:
            raise ImportProblem("The first line of that sheet has no column names.")
        rows: list[list[str]] = []
        for raw in rows_iter:
            cells = [_excel_text(cell) for cell in raw[: len(headers)]]
            if not any(cell.strip() for cell in cells):
                continue
            rows.append(cells)
            if max_rows is not None and len(rows) >= max_rows:
                break
    finally:
        book.close()
    return headers, rows


def read_sheet(
    path: Path, *, sheet: str | None = None, max_rows: int | None = None
) -> tuple[list[str], list[list[str]]]:
    """Read whichever kind of spreadsheet this is."""
    suffix = path.suffix.lower()
    if suffix in EXCEL_SUFFIXES:
        return read_xlsx(path, sheet=sheet, max_rows=max_rows)
    return read_csv(path, max_rows=max_rows)


def table_name_for(path: Path, sheet: str | None = None) -> str:
    if sheet and path.suffix.lower() in EXCEL_SUFFIXES and len(list_sheets(path)) > 1:
        return slugify_name(sheet, fallback="table")
    return slugify_name(path.stem, fallback="table")


def plan_import(file_path: str, *, sheet: str | None = None) -> dict[str, Any]:
    """Look at a spreadsheet and say what would be made, without making it.

    This is what the import wizard shows before the user presses the button.
    """
    path = resolve_under_roots(file_path, must_exist=True)
    suffix = path.suffix.lower()
    if suffix in OLD_EXCEL_SUFFIXES:
        raise ImportProblem(
            "That is an old-style Excel file (.xls). Open it in Excel or LibreOffice, "
            "save it as .xlsx, and import that."
        )
    if suffix not in CSV_SUFFIXES and suffix not in EXCEL_SUFFIXES:
        raise ImportProblem(f"Omarchy-DB does not know how to read {suffix or 'that file'} yet.")

    sheets = list_sheets(path)
    if sheet and not sheets:
        sheet = None
    headers, rows = read_sheet(path, sheet=sheet, max_rows=SAMPLE_ROWS)
    fields = infer_fields(headers, rows)
    return {
        "file": str(path),
        "table": table_name_for(path, sheet or (sheets[0] if sheets else None)),
        "sheet": sheet or (sheets[0] if sheets else ""),
        "sheets": sheets,
        "headers": headers,
        "fields": [{"name": f.name, "type": f.type, "label": f.title} for f in fields],
        "sample_rows": rows[:5],
        "rows_sampled": len(rows),
    }


def import_spreadsheet(
    storage: Storage,
    file_path: str,
    *,
    table: str | None = None,
    fields: Sequence[Field] | None = None,
    if_exists: str = "error",
    sheet: str | None = None,
) -> dict[str, Any]:
    """Read a spreadsheet (one sheet of a workbook) and put it in the database as a table."""
    path = resolve_under_roots(file_path, must_exist=True)
    plan = plan_import(str(path), sheet=sheet)
    table_name = table or plan["table"]

    if fields is None:
        fields = [Field(name=f["name"], type=f["type"], label=f["label"]) for f in plan["fields"]]
    fields = list(fields)

    headers, rows = read_sheet(path, sheet=plan["sheet"] or None)
    existed = storage.has_table(table_name)
    storage.create_table(table_name, fields, if_exists=if_exists)

    typed_rows: list[list[Any]] = []
    problems: list[str] = []
    for line_number, row in enumerate(rows, start=2):
        values: list[Any] = []
        for index, field in enumerate(fields):
            raw = row[index] if index < len(row) else None
            try:
                values.append(coerce(raw, field.type))
            except (ValueError, TypeError):
                # Never lose a cell: keep the text and note it.
                values.append(None if raw in (None, "") else str(raw))
                problems.append(f"Line {line_number}, {field.title}: kept {raw!r} as words.")
        typed_rows.append(values)

    added = storage.add_rows(table_name, fields, typed_rows)
    return {
        "file": str(path),
        "sheet": plan["sheet"],
        "table": table_name,
        "created_fields": [{"name": f.name, "type": f.type, "label": f.title} for f in fields],
        "rows_added": added,
        "notes": problems[:20],
        "note_count": len(problems),
        "replaced_existing": existed and if_exists == "replace",
    }
