"""Turn a spreadsheet into a table.

Reads CSV (and tab-separated files) and Excel `.xlsx` workbooks. Cells are
only ever read as data. In a CSV a formula is just the text it is. In an
Excel file the value Excel last worked out is used; the formula itself is
never run by Jubako.
"""

from __future__ import annotations

import csv
import datetime as _dt
from collections.abc import Sequence
from pathlib import Path
from typing import Any

from .errors import ImportProblem, JubakoError
from .fields import Field, coerce, slugify_name
from .infer import infer_fields
from .paths import resolve_under_roots
from .storage.base import Storage

SAMPLE_ROWS = 500
CSV_SUFFIXES = {".csv", ".tsv", ".txt"}
EXCEL_SUFFIXES = {".xlsx", ".xlsm"}
OLD_EXCEL_SUFFIXES = {".xls"}


def read_csv(path: Path, *, max_rows: int | None = None) -> tuple[list[str], list[list[str]]]:
    """Read a CSV file into headings plus rows of plain text."""
    with path.open("r", newline="", encoding="utf-8-sig", errors="replace") as handle:
        sample = handle.read(64 * 1024)
        handle.seek(0)
        try:
            dialect: Any = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(handle, dialect)
        try:
            headers = next(reader)
        except StopIteration:
            raise ImportProblem("That file is empty.") from None
        rows: list[list[str]] = []
        for row in reader:
            if not any(str(cell).strip() for cell in row):
                continue
            rows.append([str(cell) for cell in row])
            if max_rows is not None and len(rows) >= max_rows:
                break
    headers = [str(header) for header in headers]
    if not any(header.strip() for header in headers):
        raise ImportProblem("The first line of that file has no column names.")
    return headers, rows


def _openpyxl():
    try:
        import openpyxl  # noqa: PLC0415
    except ImportError as exc:  # pragma: no cover - depends on the machine
        raise ImportProblem(
            "Reading Excel files needs the 'openpyxl' package. "
            "Install it with: pip install 'jubako[xlsx]'"
        ) from exc
    return openpyxl


def _excel_text(value: Any) -> str:
    """Excel cells come typed. Turn each into the text the type guesser reads."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "yes" if value else "no"
    if isinstance(value, _dt.datetime):
        if value.time() == _dt.time(0, 0):
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, _dt.date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


#: How many leading rows to look at when hunting for the header row.
HEADER_SEARCH_ROWS = 30


def _load_book(path: Path):
    """Open a workbook with the values Excel last saved.

    Not read-only: read-only sheets do not expose their Excel Tables, and a
    Table is the best possible clue to where the real headers are.
    """
    openpyxl = _openpyxl()
    try:
        return openpyxl.load_workbook(path, data_only=True)
    except Exception as exc:  # openpyxl raises many kinds
        raise ImportProblem(f"That does not look like an Excel file: {exc}") from exc


def _largest_table(worksheet):
    """The Excel Table covering the most cells, or None if the sheet has none."""
    from openpyxl.utils import range_boundaries  # noqa: PLC0415

    best = None
    best_area = -1
    for table in getattr(worksheet, "tables", {}).values():
        try:
            min_col, min_row, max_col, max_row = range_boundaries(table.ref)
        except Exception:  # noqa: BLE001 - a malformed ref just means "no table"
            continue
        area = (max_col - min_col + 1) * (max_row - min_row + 1)
        if area > best_area:
            best, best_area = (table, (min_col, min_row, max_col, max_row)), area
    return best


def _header_row_index(grid: list[tuple]) -> int:
    """Where the column names are: the widest of the first rows, never a lone title.

    A titled sheet starts with "My List" in A1, maybe a subtitle, then blanks,
    then the real headers across many columns. So look at the first
    `HEADER_SEARCH_ROWS`, find the widest (most filled cells), and take the
    first row that is at least half that wide and has two or more filled
    cells. A lone title never qualifies; a header row with a blank cell or
    two still does. Only if no row has two filled cells is a one-cell row
    used, because then the sheet really is one column.
    """
    counts = [
        sum(1 for cell in row if _excel_text(cell).strip())
        for row in grid[:HEADER_SEARCH_ROWS]
    ]
    widest = max(counts, default=0)
    if widest == 0:
        raise ImportProblem("That sheet is empty.")
    wide_enough = max(2, (widest + 1) // 2)
    for index, count in enumerate(counts):
        if count >= wide_enough:
            return index
    return next(index for index, count in enumerate(counts) if count > 0)


def _sheet_region(worksheet) -> tuple[list[str], list[tuple]]:
    """Headers and data rows for one sheet: from its Excel Table if it has one,
    otherwise from the header row found by looking, with every column that
    has a heading or data below it."""
    table = _largest_table(worksheet)
    if table is not None:
        obj, (min_col, min_row, max_col, max_row) = table
        header_rows = int(getattr(obj, "headerRowCount", 1) or 0)
        cells = list(worksheet.iter_rows(
            min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col, values_only=True
        ))
        if header_rows and cells:
            headers = [_excel_text(cell).strip() for cell in cells[0]]
            data = cells[header_rows:]
        else:
            headers = [str(name) for name in (obj.column_names or [])]
            data = cells
        headers = _fill_blank_headings(headers)
        return headers, data

    grid = list(worksheet.iter_rows(values_only=True))
    if not grid:
        raise ImportProblem("That sheet is empty.")
    header_index = _header_row_index(grid)
    header_row = grid[header_index]
    data = grid[header_index + 1:]

    # The table is as wide as the last column with a heading or with data below.
    width = 0
    for cell_index, cell in enumerate(header_row):
        if _excel_text(cell).strip():
            width = cell_index + 1
    for row in data:
        for cell_index in range(len(row) - 1, width - 1, -1):
            if _excel_text(row[cell_index]).strip():
                width = cell_index + 1
                break
    headers = [_excel_text(cell).strip() for cell in header_row[:width]]
    headers += [""] * (width - len(headers))
    return _fill_blank_headings(headers, trim=False), data


def _fill_blank_headings(headers: list[str], *, trim: bool = True) -> list[str]:
    """A column with data but no heading is called "Column 3", not dropped.

    `trim` drops trailing blank headings (an Excel Table's own columns are
    always named, so there it only tidies); the sheet scan has already sized
    the table to the data, so it keeps them.
    """
    headers = list(headers)
    if trim:
        while headers and not headers[-1]:
            headers.pop()
    return [name or f"Column {index + 1}" for index, name in enumerate(headers)]


#: The last workbook read, parsed once: planning and importing every sheet
#: of one file otherwise opens it again and again.
_last_book: dict[str, Any] = {}


def _book_key(path: Path) -> tuple:
    stat = path.stat()
    return (str(path), stat.st_mtime_ns, stat.st_size)


def _parsed_book(path: Path) -> dict[str, Any]:
    key = _book_key(path)
    if _last_book.get("key") == key:
        return _last_book
    book = _load_book(path)
    try:
        sheets = list(book.sheetnames)
        regions = {name: _sheet_region_or_error(book[name]) for name in sheets}
    finally:
        book.close()
    _last_book.clear()
    _last_book.update({"key": key, "sheets": sheets, "regions": regions})
    return _last_book


def _sheet_region_or_error(worksheet):
    try:
        return _sheet_region(worksheet)
    except ImportProblem as error:
        return error


def list_sheets(path: Path) -> list[str]:
    """The sheet names in a workbook; a CSV has one unnamed sheet."""
    if path.suffix.lower() not in EXCEL_SUFFIXES:
        return []
    return list(_parsed_book(path)["sheets"])


def read_xlsx(
    path: Path, *, sheet: str | None = None, max_rows: int | None = None
) -> tuple[list[str], list[list[str]]]:
    """Read one sheet of an Excel workbook into headings plus rows of plain text.

    The headers are taken from the sheet's Excel Table when it has one (the
    largest, if several), and otherwise from the widest of the first rows, so
    a title sitting alone in A1 is never mistaken for the column names.

    `data_only=True` means a formula cell gives the value Excel last saved for
    it. Nothing is calculated here, and no macro is ever touched.
    """
    parsed = _parsed_book(path)
    if sheet:
        if sheet not in parsed["sheets"]:
            names = ", ".join(parsed["sheets"])
            raise ImportProblem(f"There is no sheet called {sheet!r}. This file has: {names}.")
    else:
        sheet = parsed["sheets"][0]
    region = parsed["regions"][sheet]
    if isinstance(region, ImportProblem):
        raise region
    headers, data = region
    if not headers:
        raise ImportProblem("That sheet has no column names.")
    rows: list[list[str]] = []
    for raw in data:
        cells = [_excel_text(cell) for cell in raw[: len(headers)]]
        cells += [""] * (len(headers) - len(cells))
        if not any(cell.strip() for cell in cells):
            continue
        rows.append(cells)
        if max_rows is not None and len(rows) >= max_rows:
            break
    return list(headers), rows


def read_sheet(
    path: Path, *, sheet: str | None = None, max_rows: int | None = None
) -> tuple[list[str], list[list[str]]]:
    """Read whichever kind of spreadsheet this is."""
    suffix = path.suffix.lower()
    if suffix in EXCEL_SUFFIXES:
        return read_xlsx(path, sheet=sheet, max_rows=max_rows)
    return read_csv(path, max_rows=max_rows)


def table_name_for(path: Path, sheet: str | None = None) -> str:
    if sheet and path.suffix.lower() in EXCEL_SUFFIXES and len(list_sheets(path)) > 1:
        return slugify_name(sheet, fallback="table")
    return slugify_name(path.stem, fallback="table")


def plan_import(file_path: str, *, sheet: str | None = None) -> dict[str, Any]:
    """Look at a spreadsheet and say what would be made, without making it.

    This is what the import wizard shows before the user presses the button.
    """
    path = resolve_under_roots(file_path, must_exist=True)
    suffix = path.suffix.lower()
    if suffix in OLD_EXCEL_SUFFIXES:
        raise ImportProblem(
            "That is an old-style Excel file (.xls). Open it in Excel or LibreOffice, "
            "save it as .xlsx, and import that."
        )
    if suffix not in CSV_SUFFIXES and suffix not in EXCEL_SUFFIXES:
        raise ImportProblem(f"Jubako does not know how to read {suffix or 'that file'} yet.")

    sheets = list_sheets(path)
    if sheet and not sheets:
        sheet = None
    headers, rows = read_sheet(path, sheet=sheet, max_rows=SAMPLE_ROWS)
    fields = infer_fields(headers, rows)
    return {
        "file": str(path),
        "table": table_name_for(path, sheet or (sheets[0] if sheets else None)),
        "sheet": sheet or (sheets[0] if sheets else ""),
        "sheets": sheets,
        "headers": headers,
        "fields": [{"name": f.name, "type": f.type, "label": f.title} for f in fields],
        "sample_rows": rows[:5],
        "rows_sampled": len(rows),
    }


def import_spreadsheet(
    storage: Storage,
    file_path: str,
    *,
    table: str | None = None,
    fields: Sequence[Field] | None = None,
    if_exists: str = "error",
    sheet: str | None = None,
) -> dict[str, Any]:
    """Read a spreadsheet (one sheet of a workbook) and put it in the database as a table."""
    path = resolve_under_roots(file_path, must_exist=True)
    plan = plan_import(str(path), sheet=sheet)
    table_name = table or plan["table"]

    if fields is None:
        fields = [Field(name=f["name"], type=f["type"], label=f["label"]) for f in plan["fields"]]
    fields = list(fields)

    headers, rows = read_sheet(path, sheet=plan["sheet"] or None)
    existed = storage.has_table(table_name)
    storage.create_table(table_name, fields, if_exists=if_exists)

    typed_rows: list[list[Any]] = []
    problems: list[str] = []
    for line_number, row in enumerate(rows, start=2):
        values: list[Any] = []
        for index, field in enumerate(fields):
            raw = row[index] if index < len(row) else None
            try:
                values.append(coerce(raw, field.type))
            except (ValueError, TypeError):
                # Never lose a cell: keep the text and note it.
                values.append(None if raw in (None, "") else str(raw))
                problems.append(f"Line {line_number}, {field.title}: kept {raw!r} as words.")
        typed_rows.append(values)

    added = storage.add_rows(table_name, fields, typed_rows)
    return {
        "file": str(path),
        "sheet": plan["sheet"],
        "table": table_name,
        "created_fields": [{"name": f.name, "type": f.type, "label": f.title} for f in fields],
        "rows_added": added,
        "notes": problems[:20],
        "note_count": len(problems),
        "replaced_existing": existed and if_exists == "replace",
    }


def import_workbook(
    storage: Storage,
    file_path: str,
    *,
    if_exists: str = "error",
) -> dict[str, Any]:
    """Every sheet of a workbook becomes its own table, named after the sheet.

    Each sheet uses the guessed types. A sheet that cannot be read (empty, or
    no column names) is reported, not fatal; the others still go in. With
    `if_exists="error"` a sheet whose table already exists is reported the
    same way. A CSV is one sheet, so it just imports.
    """
    path = resolve_under_roots(file_path, must_exist=True)
    sheets = list_sheets(path) or [None]
    tables: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    for sheet in sheets:
        try:
            result = import_spreadsheet(storage, str(path), sheet=sheet, if_exists=if_exists)
        except JubakoError as error:
            errors.append({"sheet": sheet or path.name, "error": str(error)})
            continue
        tables.append({
            "sheet": result["sheet"], "table": result["table"], "rows_added": result["rows_added"],
            "fields": len(result["created_fields"]), "note_count": result["note_count"],
            "replaced_existing": result["replaced_existing"],
        })
    return {"file": str(path), "tables": tables, "errors": errors}


def workbook_plan(file_path: str) -> list[dict[str, Any]]:
    """For each sheet: the table it would become, and whether that table exists is up to the caller."""
    path = resolve_under_roots(file_path, must_exist=True)
    sheets = list_sheets(path) or [None]
    out = []
    for sheet in sheets:
        out.append({"sheet": sheet or "", "table": table_name_for(path, sheet)})
    return out
