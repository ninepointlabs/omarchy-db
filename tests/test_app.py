"""The QML desktop app: the bridge, the rows model, the theme, and the window itself.

These run without a display (`QT_QPA_PLATFORM=offscreen`). They are skipped
when PySide6 is not installed, because the core library and the MCP server do
not need it.
"""

from __future__ import annotations

import os
from pathlib import Path

import pytest

PySide6 = pytest.importorskip("PySide6")

from PySide6.QtCore import QObject, QUrl  # noqa: E402

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

QML_DIR = Path(__file__).resolve().parents[1] / "src" / "jubako_app" / "qml"


@pytest.fixture(scope="session")
def app():
    from PySide6.QtGui import QGuiApplication

    from jubako_app.main import make_app

    return QGuiApplication.instance() or make_app(["jubako-app"])


def wait_for_import(bridge, timeout_ms: int = 30000) -> dict:
    """Run the event loop until the worker reports back (polling would starve it of the GIL)."""
    from PySide6.QtCore import QEventLoop, QTimer

    loop = QEventLoop()
    results = []

    def done(result):
        results.append(dict(result))
        loop.quit()

    bridge.importFinished.connect(done)
    try:
        QTimer.singleShot(timeout_ms, loop.quit)
        loop.exec()
    finally:
        bridge.importFinished.disconnect(done)
    assert results, "the import never finished"
    return results[0]


@pytest.fixture
def bridge(app, sandbox):
    from jubako_app.bridge import Bridge

    bridge = Bridge()
    yield bridge
    bridge.closeDatabase()


# -- the bridge ------------------------------------------------------------

def test_backends_lists_the_three_engines_sqlite_first(bridge):
    keys = [b["key"] for b in bridge.backends()]
    assert keys == ["sqlite", "postgres", "mysql"]
    assert bridge.backends()[0]["ready"] is True
    assert bridge.backends()[0]["needsServer"] is False


def test_new_database_then_import_then_rows(bridge, sandbox, pets_csv):
    assert bridge.isOpen is False
    made = bridge.newDatabase("sqlite", str(sandbox / "pets.jubadb"), "My Pets", {})
    assert made == {"ok": True}
    assert bridge.isOpen is True
    assert bridge.title == "My Pets"
    assert bridge.tables == []

    done = bridge.importSpreadsheet(str(pets_csv), False)
    assert done["ok"] is True
    assert done["table"] == "pets"
    assert done["rowsAdded"] == 4
    assert [t["name"] for t in bridge.tables] == ["pets"]
    assert bridge.currentTable == "pets"
    assert bridge.totalRows == 4
    assert bridge.rows.rowCount() == 4
    assert bridge.rows.columnCount() == 6  # id + five fields


def test_import_asks_before_replacing_a_table(bridge, sandbox, pets_csv):
    bridge.newDatabase("sqlite", str(sandbox / "pets.jubadb"), "", {})
    bridge.importSpreadsheet(str(pets_csv), False)
    again = bridge.importSpreadsheet(str(pets_csv), False)
    assert again["ok"] is False
    assert again["needsConfirm"] is True
    assert again["table"] == "pets"
    assert bridge.totalRows == 4
    replaced = bridge.importSpreadsheet(str(pets_csv), True)
    assert replaced["ok"] is True
    assert bridge.totalRows == 4


def test_file_urls_are_accepted(bridge, sandbox, pets_csv):
    url = QUrl.fromLocalFile(str(sandbox / "from-url.jubadb")).toString()
    assert bridge.newDatabase("sqlite", url, "", {})["ok"] is True
    assert bridge.location == str(sandbox / "from-url.jubadb")
    csv_url = QUrl.fromLocalFile(str(pets_csv)).toString()
    assert bridge.importSpreadsheet(csv_url, False)["ok"] is True
    assert bridge.suggestedDatabaseName(csv_url) == "pets.jubadb"


def test_import_into_new_makes_the_database_and_names_it(bridge, sandbox, pets_csv):
    result = bridge.importIntoNew(str(pets_csv), str(sandbox / "new.jubadb"))
    assert result["ok"] is True
    assert bridge.title == "Pets"
    assert bridge.totalRows == 4


def test_paths_outside_the_approved_roots_are_refused(bridge, sandbox):
    result = bridge.newDatabase("sqlite", "/etc/evil.jubadb", "", {})
    assert result["ok"] is False
    assert "outside" in result["error"]
    assert bridge.isOpen is False


def test_recent_and_open_recent(bridge, sandbox):
    bridge.newDatabase("sqlite", str(sandbox / "a.jubadb"), "A", {})
    bridge.newDatabase("sqlite", str(sandbox / "b.jubadb"), "B", {})
    bridge.closeDatabase()
    recent = bridge.recent()
    assert [r["title"] for r in recent] == ["B", "A"]
    assert recent[0]["backendWord"] == "File"
    assert bridge.openRecent(1) == {"ok": True}
    assert bridge.title == "A"


def test_open_recent_forgets_a_database_that_is_gone(bridge, sandbox):
    bridge.newDatabase("sqlite", str(sandbox / "gone.jubadb"), "Gone", {})
    bridge.closeDatabase()
    (sandbox / "gone.jubadb").unlink()
    assert bridge.openRecent(0)["ok"] is False
    assert bridge.recent() == []


def test_a_remembered_server_asks_for_its_password(bridge, sandbox):
    from jubako import catalog

    catalog.remember(
        title="Shop", backend="postgres",
        where="host=db.local port=5433 dbname=shop user=tim password=***",
    )
    result = bridge.openRecent(0)
    assert result["ok"] is False
    assert result["needsConnection"] is True
    assert result["backend"] == "postgres"
    assert result["connection"] == {"host": "db.local", "port": "5433", "database": "shop", "user": "tim"}

    catalog.remember(title="Shop2", backend="mysql", where="root@127.0.0.1:3306/shop")
    result = bridge.openRecent(0)
    assert result["connection"] == {"host": "127.0.0.1", "port": "3306", "database": "shop", "user": "root"}


def test_server_backend_errors_come_back_as_words(bridge):
    result = bridge.newDatabase("postgres", "", "X", {"host": "", "port": ""})
    assert result["ok"] is False
    assert result["error"]  # either "needs psycopg" or "which database to use"
    assert bridge.isOpen is False


# -- the rows model -----------------------------------------------------------

def test_rows_model_shows_yes_no_and_blanks(app):
    from PySide6.QtCore import Qt

    from jubako_app.bridge import RowsModel

    model = RowsModel()
    model.load(
        {
            "fields": [
                {"name": "name", "type": "text", "label": "Name"},
                {"name": "good", "type": "boolean", "label": "Good?"},
            ],
            "rows": [[1, "Rex", True], [2, None, False]],
        }
    )
    assert model.columnCount() == 3
    assert model.headerData(1, Qt.Orientation.Horizontal) == "Name"
    assert model.data(model.index(0, 2)) == "Yes"
    assert model.data(model.index(1, 2)) == "No"
    assert model.data(model.index(1, 1)) == ""
    assert model.columnWidth(1) >= 60
    assert model.kindAt(2) == "boolean"


# -- the theme ----------------------------------------------------------------

def test_theme_reads_omarchy_colors_and_falls_back(tmp_path):
    from PySide6.QtGui import QPalette

    from jubako_app.theme import build_palette, load_colors

    missing = load_colors(tmp_path / "nope.toml")
    assert missing["accent"]  # the fallback
    file = tmp_path / "colors.toml"
    file.write_text('mode = "light"\naccent = "#123456"\nbackground = "#ffffff"\n')
    colors = load_colors(file)
    assert colors["accent"] == "#123456"
    assert colors["mode"] == "light"
    palette = build_palette(colors)
    assert palette.color(QPalette.ColorRole.Highlight).name() == "#123456"
    assert palette.color(QPalette.ColorRole.Base).name() == "#ffffff"


# -- the window ----------------------------------------------------------------

def test_the_window_loads_and_shows_an_imported_table(app, bridge, sandbox, pets_csv):
    from PySide6.QtQml import QQmlApplicationEngine

    from jubako_app.theme import Theme

    from jubako_app.report_bridge import Report, ReportImageProvider

    engine = QQmlApplicationEngine()
    provider = ReportImageProvider()
    engine.addImageProvider("report", provider)
    engine.rootContext().setContextProperty("Bridge", bridge)
    engine.rootContext().setContextProperty("Report", Report(bridge.storage, provider))
    engine.rootContext().setContextProperty("Theme", Theme(app))
    engine.load(QUrl.fromLocalFile(str(QML_DIR / "Main.qml")))
    assert len(engine.rootObjects()) == 1, "Main.qml did not load"
    window = engine.rootObjects()[0]
    assert window.title() == "Jubako"

    assert bridge.importIntoNew(str(pets_csv), str(sandbox / "pets.jubadb"))["ok"] is True
    app.processEvents()
    assert window.title() == "Pets — Jubako"
    grids = [
        child for child in window.findChildren(QObject)
        if child.metaObject().className() == "QQuickTableView"
    ]
    assert grids, "the rows grid is not on screen"
    assert grids[0].property("rows") == 4
    assert grids[0].property("columns") == 6

    bridge.closeDatabase()
    app.processEvents()
    assert window.title() == "Jubako"
    engine.deleteLater()


def test_theme_follows_omarchy_theme_set_swapping_the_directory(app, tmp_path, monkeypatch):
    """`omarchy theme set` does rm -rf current/theme; mv theme.next current/theme."""
    import shutil

    from PySide6.QtTest import QTest

    from jubako_app.theme import Theme

    monkeypatch.delenv("JUBAKO_THEME_FILE", raising=False)
    monkeypatch.setenv("XDG_STATE_HOME", str(tmp_path))
    current = tmp_path / "omarchy" / "current"
    (current / "theme").mkdir(parents=True)
    (current / "theme" / "colors.toml").write_text('mode = "dark"\naccent = "#111111"\n')
    (current / "theme.name").write_text("one\n")

    theme = Theme()
    assert theme.accent == "#111111"
    seen = []
    theme.changed.connect(lambda: seen.append(theme.accent))

    staging = current / "theme.next"
    staging.mkdir()
    (staging / "colors.toml").write_text('mode = "light"\naccent = "#222222"\n')
    shutil.rmtree(current / "theme")
    staging.rename(current / "theme")
    (current / "theme.name").write_text("two\n")

    QTest.qWait(700)
    assert seen == ["#222222"]
    assert theme.isDark is False

    # And again, to prove the watches were re-armed after the first swap.
    staging.mkdir()
    (staging / "colors.toml").write_text('mode = "dark"\naccent = "#333333"\n')
    shutil.rmtree(current / "theme")
    staging.rename(current / "theme")
    QTest.qWait(700)
    assert seen == ["#222222", "#333333"]
    assert str(current / "theme" / "colors.toml") in theme.watched()


# -- editing rows -------------------------------------------------------------

@pytest.fixture
def pets(bridge, sandbox, pets_csv):
    assert bridge.importIntoNew(str(pets_csv), str(sandbox / "pets.jubadb"))["ok"] is True
    return bridge


def test_save_row_adds_and_updates(pets):
    assert [f["name"] for f in pets.fields] == ["name", "age", "adopted_on", "is_good", "weight_kg"]
    added = pets.saveRow(0, {"name": "Nova", "age": "3", "adopted_on": "2025-02-01", "is_good": True, "weight_kg": ""})
    assert added["ok"] is True
    assert added["rowId"] > 0
    assert pets.totalRows == 5
    assert added["rowIndex"] == 4
    record = pets.rows.record(added["rowIndex"])
    assert record["id"] == added["rowId"]
    assert record["values"] == {"name": "Nova", "age": "3", "adopted_on": "2025-02-01", "is_good": True, "weight_kg": ""}

    changed = pets.saveRow(added["rowId"], {"age": 4, "is_good": False})
    assert changed["ok"] is True
    assert pets.rows.record(4)["values"]["age"] == "4"
    assert pets.rows.record(4)["values"]["is_good"] is False
    assert pets.totalRows == 5


def test_save_row_explains_a_value_that_does_not_fit(pets):
    result = pets.saveRow(0, {"name": "Bad", "age": "three"})
    assert result["ok"] is False
    assert result["error"] == "age needs a whole number. “three” does not fit."
    assert pets.totalRows == 4
    result = pets.saveRow(pets.rows.rowId(0), {"adopted_on": "someday"})
    assert result["ok"] is False
    assert "needs a date" in result["error"]


def test_delete_row(pets):
    first = pets.rows.rowId(0)
    assert pets.deleteRow(first)["ok"] is True
    assert pets.totalRows == 3
    assert pets.rows.rowIndexOf(first) == -1
    again = pets.deleteRow(first)
    assert again["ok"] is False
    assert pets.totalRows == 3


def test_grid_edit_goes_through_set_data(pets):
    from PySide6.QtCore import Qt

    model = pets.rows
    assert model.flags(model.index(0, 0)) & Qt.ItemFlag.ItemIsEditable == Qt.ItemFlag(0)
    assert model.flags(model.index(0, 1)) & Qt.ItemFlag.ItemIsEditable
    assert model.setData(model.index(0, 1), "Rexy") is True
    assert model.data(model.index(0, 1)) == "Rexy"
    assert model.setData(model.index(0, 4), "no") is True
    assert model.data(model.index(0, 4)) == "No"
    heard = []
    pets.message.connect(heard.append)
    assert model.setData(model.index(0, 2), "lots") is False
    assert heard and "whole number" in heard[0]
    assert model.setData(model.index(0, 0), "7") is False


# -- the import wizard and the worker thread ------------------------------------

def test_plan_import_and_import_planned_off_the_gui_thread(bridge, sandbox, pets_csv):
    bridge.newDatabase("sqlite", str(sandbox / "w.jubadb"), "W", {})
    plan = bridge.planImport(str(pets_csv), "")
    assert plan["ok"] is True
    assert plan["table"] == "pets"
    assert plan["exists"] is False
    assert [f["type"] for f in plan["fields"]][:2] == ["text", "integer"]

    # Change a guess: keep age as words, and call the table something else.
    fields = [dict(f) for f in plan["fields"]]
    fields[1]["type"] = "text"
    started = bridge.importPlanned(str(pets_csv), "", "animals", fields, False)
    assert started == {"ok": True, "pending": True}
    assert bridge.busy is True
    result = wait_for_import(bridge)
    assert result["ok"] is True, result
    assert result["table"] == "animals"
    assert result["rowsAdded"] == 4
    assert bridge.busy is False
    assert bridge.currentTable == "animals"
    assert bridge.fields[1]["type"] == "text"
    assert bridge.rows.record(0)["values"]["age"] == "4"

    # A second import of the same table asks first, synchronously.
    again = bridge.importPlanned(str(pets_csv), "", "animals", [], False)
    assert again["needsConfirm"] is True


def test_export_table_from_the_window(bridge, sandbox, pets_csv):
    pytest.importorskip("openpyxl")
    bridge.importIntoNew(str(pets_csv), str(sandbox / "e.jubadb"))
    out = bridge.exportTable(QUrl.fromLocalFile(str(sandbox / "out.xlsx")).toString(), "xlsx")
    assert out["ok"] is True
    assert out["rows"] == 4
    assert (sandbox / "out.xlsx").exists()
    assert bridge.exportTable("/etc/out.csv", "csv")["ok"] is False


def test_report_bridge_builds_previews_and_keeps(app, bridge, sandbox, pets_csv):
    from jubako_app.report_bridge import Report, ReportImageProvider

    bridge.importIntoNew(str(pets_csv), str(sandbox / "r.jubadb"))
    provider = ReportImageProvider()
    report = Report(bridge.storage, provider)
    built = report.build({"table": "pets", "title": "All pets", "columns": ["name", "age"], "margins_mm": "20"})
    assert built["ok"] is True
    assert built["pages"] == 1
    assert report.spec["columns"] == ["name", "age"]
    assert report.previewSource(0).startswith("image://report/0?")
    from PySide6.QtCore import QSize

    image = provider.requestImage("0?1", QSize(), QSize(400, 0))
    assert image.width() == 400
    assert image.height() > image.width()  # portrait

    saved = report.savePdf(QUrl.fromLocalFile(str(sandbox / "pets.pdf")).toString())
    assert saved["ok"] is True
    assert (sandbox / "pets.pdf").read_bytes().startswith(b"%PDF")

    assert report.keep("Names and ages")["ok"] is True
    assert [r["name"] for r in report.kept()] == ["Names and ages"]
    report.build({"table": "pets"})
    loaded = report.load("Names and ages")
    assert loaded["ok"] is True
    assert loaded["spec"]["columns"] == ["name", "age"]
    assert report.forget("Names and ages")["ok"] is True
    assert report.build({"table": "pets", "columns": ["nope"]})["ok"] is False
    assert "nope" in report.error


# -- fields, tables, the whole file ------------------------------------------------

def test_rename_and_delete_field_from_the_bridge(pets):
    renamed = pets.renameField("weight_kg", "weight", "Weight")
    assert renamed["ok"] is True
    assert pets.fields[-1] == {"name": "weight", "type": "real", "label": "Weight"}
    assert pets.rows.headerData(5, PySide6.QtCore.Qt.Orientation.Horizontal) == "Weight"
    relabel = pets.renameField("age", "", "Years")
    assert relabel["ok"] is True
    assert pets.fields[1]["label"] == "Years"
    assert pets.renameField("age", "name", "")["ok"] is False
    gone = pets.deleteField("adopted_on")
    assert gone["ok"] is True
    assert [f["name"] for f in pets.fields] == ["name", "age", "is_good", "weight"]
    assert pets.rows.columnCount() == 5
    assert pets.totalRows == 4


def test_delete_table_from_the_bridge(pets):
    assert pets.deleteTable("pets") == {"ok": True, "rowsDeleted": 4}
    assert pets.tables == []
    assert pets.currentTable == ""
    assert pets.deleteTable("pets")["ok"] is False


def test_delete_database_from_the_bridge(bridge, sandbox, pets_csv):
    bridge.importIntoNew(str(pets_csv), str(sandbox / "doomed.jubadb"))
    assert bridge.isLocalFile is True
    path = bridge.location
    result = bridge.deleteDatabase()
    assert result == {"ok": True, "file": path}
    assert bridge.isOpen is False
    assert not (sandbox / "doomed.jubadb").exists()
    assert bridge.recent() == []
    assert bridge.deleteDatabase()["ok"] is False


def test_recent_list_remove_and_delete_file(bridge, sandbox):
    bridge.newDatabase("sqlite", str(sandbox / "a.jubadb"), "A", {})
    bridge.newDatabase("sqlite", str(sandbox / "b.jubadb"), "B", {})
    bridge.closeDatabase()
    assert [r["title"] for r in bridge.recent()] == ["B", "A"]
    assert bridge.removeRecent(0)["ok"] is True
    assert [r["title"] for r in bridge.recent()] == ["A"]
    assert (sandbox / "b.jubadb").exists()
    assert bridge.deleteRecentFile(0)["ok"] is True
    assert not (sandbox / "a.jubadb").exists()
    assert bridge.recent() == []


def test_import_every_sheet_from_the_bridge(bridge, sandbox):
    openpyxl = pytest.importorskip("openpyxl")

    book = openpyxl.Workbook()
    a = book.active; a.title = "People"; a.append(["Name", "Age"]); a.append(["Ann", 31])
    b = book.create_sheet("Places"); b.append(["City"]); b.append(["Tyler"])
    book.create_sheet("Blank")
    book.save(sandbox / "three.xlsx")
    bridge.newDatabase("sqlite", str(sandbox / "all.jubadb"), "All", {})
    url = QUrl.fromLocalFile(str(sandbox / "three.xlsx")).toString()

    plan = bridge.workbookPlan(url)
    assert [(p["sheet"], p["table"], p["exists"]) for p in plan] == [
        ("People", "people", False), ("Places", "places", False), ("Blank", "blank", False)]

    assert bridge.importAllSheets(url, False) == {"ok": True, "pending": True}
    result = wait_for_import(bridge)
    assert result["ok"] is True, result
    assert result["tables"] == ["people", "places"]
    assert len(result["errors"]) == 1 and result["errors"][0].startswith("Blank:")
    assert [t["name"] for t in bridge.tables] == ["people", "places"]

    # A second time, one confirm covers every clash.
    again = bridge.importAllSheets(url, False)
    assert again["needsConfirm"] is True
    assert again["tables"] == ["people", "places"]


# -- filters and add field -----------------------------------------------------------

def test_add_field_and_filter_from_the_bridge(pets):
    assert pets.slugName("Vet's phone") == "vet_s_phone"
    assert pets.slugName("  ") == ""
    added = pets.addField("Moved", "", "boolean")
    assert added["ok"] is True
    assert pets.fields[-1] == {"name": "moved", "type": "boolean", "label": "Moved"}
    assert pets.rows.columnCount() == 7
    assert pets.rows.record(0)["values"]["moved"] is False
    assert pets.addField("Moved", "", "text")["ok"] is False
    assert pets.addField("", "", "text")["ok"] is False

    first = pets.rows.rowId(0)
    third = pets.rows.rowId(2)
    assert pets.saveRow(first, {"moved": True})["ok"] is True
    assert pets.saveRow(third, {"moved": "yes"})["ok"] is True

    ops = [o["key"] for o in pets.filterOps()]
    assert ops == ["is", "is_not", "empty", "not_empty", "contains"]
    shown = pets.setFilter("moved", "is_not", "yes")
    assert shown["ok"] is True
    assert pets.totalRows == 2
    assert pets.allRows == 4
    assert pets.filterWords == "Moved is not Yes"
    assert pets.filter == {"field": "moved", "op": "is_not", "value": "yes"}
    assert [pets.rows.record(i)["values"]["name"] for i in range(pets.rows.rowCount())] == ["Milo", "Pip"]

    # Editing keeps the filter; a new row that does not match drops out of view but is saved.
    assert pets.saveRow(0, {"name": "Nova", "moved": True})["ok"] is True
    assert pets.totalRows == 2 and pets.allRows == 5
    assert pets.saveRow(0, {"name": "Ash"})["ok"] is True
    assert pets.totalRows == 3
    assert pets.filterWords == "Moved is not Yes"

    bad = pets.setFilter("age", "is", "old")
    assert bad["ok"] is False and "whole number" in bad["error"]
    assert pets.filterWords == "Moved is not Yes"  # the old filter stays
    assert pets.clearFilter()["ok"] is True
    assert pets.filterWords == "" and pets.totalRows == 6 and pets.filter == {}

    # The filter is remembered per table while the database is open, and dropped if its field goes.
    pets.setFilter("moved", "empty", "")
    pets.selectTable("pets")
    assert pets.filterWords == "Moved is empty"
    pets.deleteField("moved")
    assert pets.filterWords == "" and pets.totalRows == 6


def test_report_bridge_uses_the_table_filter(app, bridge, sandbox, pets_csv):
    from jubako_app.report_bridge import Report, ReportImageProvider

    bridge.importIntoNew(str(pets_csv), str(sandbox / "f.jubadb"))
    bridge.setFilter("is_good", "is", "no")
    report = Report(bridge.storage, ReportImageProvider())
    built = report.build({"table": "pets", "columns": ["name"], "filter": bridge.filter})
    assert built["ok"] is True and built["rows"] == 1
    assert report.spec["filter_words"] == "is_good is No"
    assert report.build({"table": "pets", "columns": ["name"], "filter": None})["rows"] == 4


# -- saved views ---------------------------------------------------------------------

def test_saved_views_from_the_bridge(pets, sandbox):
    pets.addField("Moved", "", "boolean")
    pets.saveRow(pets.rows.rowId(0), {"moved": "yes"})
    assert pets.views == []
    assert pets.saveView("Still here", False, False)["ok"] is False  # no filter yet

    pets.setFilter("moved", "is_not", "yes")
    saved = pets.saveView("Still here", False, False)
    assert saved == {"ok": True, "name": "Still here"}
    assert pets.currentView == "Still here"
    assert pets.views == [{"name": "Still here", "words": "Moved is not Yes", "default": False}]

    # Editing the bar detaches from the view; saving under the same name asks first.
    pets.setFilter("moved", "empty", "")
    assert pets.currentView == ""
    again = pets.saveView("Still here", False, False)
    assert again["needsConfirm"] is True and again["name"] == "Still here"
    assert pets.saveView("Still here", True, True)["ok"] is True
    assert pets.views[0]["words"] == "Moved is empty" and pets.views[0]["default"] is True

    pets.clearFilter()
    assert pets.totalRows == 4 and pets.currentView == ""
    applied = pets.applyView("Still here")
    assert applied["ok"] is True
    assert pets.totalRows == 3 and pets.filterWords == "Moved is empty" and pets.currentView == "Still here"

    assert pets.renameView("Still here", "Not moved")["ok"] is True
    assert pets.currentView == "Not moved"
    assert pets.setDefaultView("Not moved", False)["ok"] is True
    assert pets.views[0]["default"] is False

    # It lives in the file: close, reopen, and it is there.
    path = pets.location
    pets.closeDatabase()
    assert pets.openDatabase("sqlite", path, {})["ok"] is True
    assert [v["name"] for v in pets.views] == ["Not moved"]
    assert pets.applyView("Not moved")["ok"] is True and pets.totalRows == 3

    # Deleting the field the view is on takes the view with it, quietly.
    pets.clearFilter()
    pets.deleteField("moved")
    assert pets.views == []
    result = pets.applyView("Not moved")
    assert result["ok"] is False and "no view" in result["error"]
    assert pets.deleteView("Not moved")["ok"] is False


def test_a_default_view_opens_with_the_table(bridge, sandbox, pets_csv):
    bridge.importIntoNew(str(pets_csv), str(sandbox / "d.jubadb"))
    bridge.setFilter("is_good", "is", "yes")
    assert bridge.saveView("Good ones", False, True)["ok"] is True
    path = bridge.location
    bridge.closeDatabase()
    bridge.openDatabase("sqlite", path, {})
    assert bridge.currentView == "Good ones"
    assert bridge.totalRows == 3 and bridge.allRows == 4
    bridge.clearFilter()
    bridge.selectTable("pets")
    assert bridge.totalRows == 4  # the default is offered once per session, not forced
    assert bridge.deleteView("Good ones")["ok"] is True
    assert bridge.views == []


# -- the Add field dialog, through the QML itself ------------------------------------------

def test_add_field_through_the_window(app, bridge, sandbox, pets_csv):
    """The path Tim uses: More… → Add field…, type a label, pick yes/no, Add it."""
    from PySide6.QtCore import Q_ARG, QMetaObject
    from PySide6.QtQml import QQmlApplicationEngine

    from jubako_app.report_bridge import Report, ReportImageProvider
    from jubako_app.theme import Theme

    engine = QQmlApplicationEngine()
    provider = ReportImageProvider()
    engine.addImageProvider("report", provider)
    engine.rootContext().setContextProperty("Bridge", bridge)
    engine.rootContext().setContextProperty("Report", Report(bridge.storage, provider))
    engine.rootContext().setContextProperty("Theme", Theme(app))
    engine.load(QUrl.fromLocalFile(str(QML_DIR / "Main.qml")))
    window = engine.rootObjects()[0]
    assert bridge.importIntoNew(str(pets_csv), str(sandbox / "ui.jubadb"))["ok"] is True
    app.processEvents()

    dialog = window.findChild(QObject, "addFieldDialog")
    assert dialog is not None, "the Add field dialog is not at page level"
    assert QMetaObject.invokeMethod(dialog, "openFor", Q_ARG("QVariant", False))
    app.processEvents()
    assert dialog.property("visible") is True
    label = dialog.findChild(QObject, "addFieldLabel")
    name = dialog.findChild(QObject, "addFieldName")
    kind = dialog.findChild(QObject, "addFieldType")
    label.setProperty("text", "Moved")
    QMetaObject.invokeMethod(label, "textEdited")     # what typing does: the name follows
    app.processEvents()
    assert name.property("text") == "moved"
    types = [t["key"] for t in bridge.fieldTypes()]
    kind.setProperty("currentIndex", types.index("boolean"))
    go = dialog.findChild(QObject, "addFieldGo")
    assert go.property("enabled") is True
    QMetaObject.invokeMethod(go, "clicked")
    app.processEvents()

    assert dialog.property("visible") is False
    assert bridge.fields[-1] == {"name": "moved", "type": "boolean", "label": "Moved"}
    assert bridge.rows.columnCount() == 7
    assert [bridge.rows.record(i)["values"]["moved"] for i in range(4)] == [False] * 4
    grid = [c for c in window.findChildren(QObject) if c.metaObject().className() == "QQuickTableView"][0]
    assert grid.property("columns") == 7

    # From the Fields list: Add field… closes the list, and the list comes back afterwards.
    fields_dialog = window.findChild(QObject, "fieldsDialog")
    QMetaObject.invokeMethod(fields_dialog, "openFor")
    app.processEvents()
    add_button = fields_dialog.findChild(QObject, "fieldsAddButton")
    QMetaObject.invokeMethod(add_button, "clicked")
    app.processEvents()
    assert fields_dialog.property("visible") is False
    assert dialog.property("visible") is True
    label.setProperty("text", "Notes")
    QMetaObject.invokeMethod(label, "textEdited")
    QMetaObject.invokeMethod(go, "clicked")
    app.processEvents()
    assert dialog.property("visible") is False
    assert fields_dialog.property("visible") is True
    assert [f["name"] for f in bridge.fields][-2:] == ["moved", "notes"]

    # A duplicate is refused in words and the dialog stays open.
    QMetaObject.invokeMethod(fields_dialog, "close")
    QMetaObject.invokeMethod(dialog, "openFor", Q_ARG("QVariant", False))
    label.setProperty("text", "Moved")
    QMetaObject.invokeMethod(label, "textEdited")
    QMetaObject.invokeMethod(go, "clicked")
    app.processEvents()
    assert dialog.property("visible") is True
    assert "already a field" in dialog.property("error")
    QMetaObject.invokeMethod(dialog, "close")
    engine.deleteLater()


def _load_window(app, bridge):
    from PySide6.QtQml import QQmlApplicationEngine

    from jubako_app.report_bridge import Report, ReportImageProvider
    from jubako_app.theme import Theme

    engine = QQmlApplicationEngine()
    provider = ReportImageProvider()
    engine.addImageProvider("report", provider)
    engine.rootContext().setContextProperty("Bridge", bridge)
    engine.rootContext().setContextProperty("Report", Report(bridge.storage, provider))
    theme = Theme(app)
    engine.rootContext().setContextProperty("Theme", theme)
    engine.load(QUrl.fromLocalFile(str(QML_DIR / "Main.qml")))
    return engine, engine.rootObjects()[0], theme


def test_the_window_palette_follows_a_light_theme(app, bridge, tmp_path, monkeypatch):
    """Fusion paints dialogs and plain labels from the window palette, which must be
    the Omarchy theme, not Fusion's stock palette for the system colour scheme."""
    from PySide6.QtCore import Qt
    from PySide6.QtQml import QQmlEngine, QQmlExpression

    file = tmp_path / "colors.toml"
    file.write_text(
        'mode = "light"\nbackground = "#eff1f5"\ndark_background = "#e3e4e8"\n'
        'foreground = "#4c4f69"\naccent = "#1e66f5"\n'
    )
    monkeypatch.setenv("JUBAKO_THEME_FILE", str(file))
    engine, window, _ = _load_window(app, bridge)
    try:
        assert app.styleHints().colorScheme() == Qt.ColorScheme.Light
        expr = QQmlExpression(
            QQmlEngine.contextForObject(window), window,
            "[win.palette.windowText, win.palette.text, win.palette.base, win.palette.window].join(' ')",
        )
        assert expr.evaluate()[0] == "#4c4f69 #4c4f69 #eff1f5 #e3e4e8"
    finally:
        engine.deleteLater()
        monkeypatch.delenv("JUBAKO_THEME_FILE")
        from jubako_app.theme import Theme

        Theme(app)  # put the session's palette back


def test_the_filter_bar_starts_fresh_on_another_table(app, bridge, sandbox):
    import openpyxl
    from PySide6.QtCore import QEvent

    book = openpyxl.Workbook()
    book.active.title = "People"
    book.active.append(["Name", "Moved"])
    book.active.append(["Ann", True])
    places = book.create_sheet("Places")
    places.append(["Town", "Visited"])
    places.append(["Bend", False])
    xlsx = sandbox / "two.xlsx"
    book.save(xlsx)
    engine, window, _ = _load_window(app, bridge)
    assert bridge.newDatabase("sqlite", str(sandbox / "two.jubadb"), "Two", {})["ok"] is True
    assert bridge.importAllSheets(str(xlsx), False)["ok"] is True
    wait_for_import(bridge)
    try:
        app.processEvents()
        value = window.findChild(QObject, "filterValue")
        bridge.selectTable("people")
        assert bridge.setFilter("moved", "is_not", "yes")["ok"] is True
        app.processEvents()
        assert value.property("text") == "yes"
        bridge.clearFilter()
        value.setProperty("text", "typed but not applied")
        bridge.selectTable("places")
        app.processEvents()
        assert value.property("text") == ""
    finally:
        engine.deleteLater()
        app.sendPostedEvents(None, QEvent.Type.DeferredDelete)


def test_the_views_menu_is_wide_enough_for_its_rows(app, bridge, sandbox, pets_csv):
    from PySide6.QtCore import QEvent
    from PySide6.QtQml import QQmlEngine, QQmlExpression

    engine, window, _ = _load_window(app, bridge)
    window.setProperty("width", 1080)
    assert bridge.importIntoNew(str(pets_csv), str(sandbox / "menu.jubadb"))["ok"] is True
    assert bridge.setFilter("name", "contains", "i")["ok"] is True
    assert bridge.saveView("A view with quite a long name to widen the menu", False, False)["ok"] is True
    try:
        app.processEvents()
        button = window.findChild(QObject, "viewsButton")
        menu = window.findChild(QObject, "viewsMenu")

        def ask(code):  # evaluated where viewsMenu is in scope
            return QQmlExpression(QQmlEngine.contextForObject(button), button, code).evaluate()[0]

        ask("viewsMenu.open()")
        app.processEvents()
        assert menu.property("visible") is True
        # Its right edge stays inside the window, and the saved view's row is not cut short.
        assert ask("viewsMenu.contentItem.mapToItem(null, viewsMenu.width, 0).x") <= window.property("width") + 1
        row = ("(function(){ for (let i = 0; i < viewsMenu.count; i++) { const it = viewsMenu.itemAt(i);"
               " if (it && String(it.text).startsWith('A view with')) return it } })()")
        assert ask(f"{row}.contentItem.implicitWidth <= {row}.availableWidth + 1") is True
    finally:
        engine.deleteLater()
        app.sendPostedEvents(None, QEvent.Type.DeferredDelete)
